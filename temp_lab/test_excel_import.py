"""Test Excel upload API with adequate timeouts"""
import httpx, json, asyncio, io

async def test_excel():
    base = 'http://127.0.0.1:8000/api/ai'
    from openpyxl import Workbook

    async def do_test(name, rows, sheet_type="position"):
        print(f"\n=== {name} ===")
        async with httpx.AsyncClient(timeout=120.0) as c:
            r = await c.post(f'{base}/upload-excel',
                files={'file': ('test.xlsx', buf, 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')},
                data={'sheet_type': sheet_type})
            d = r.json()
            print(f'  status={r.status_code} count={d.get("count",0)}')
            if d.get("data"):
                for item in d["data"][:3]:
                    print(f'  stock_code={item.get("stock_code")} name={item.get("stock_name")} qty={item.get("shares")} cost={item.get("cost_price")}')
            elif d.get("message"):
                print(f'  msg={d["message"][:100]}')
            assert r.status_code == 200
            return d

    # 1. 标准中文列名
    wb = Workbook(); ws = wb.active
    ws.append(["证券代码", "证券名称", "持股数量", "成本价"])
    ws.append(["600519", "贵州茅台", 100, 1750.00])
    ws.append(["000858", "五粮液", 200, 135.00])
    buf = io.BytesIO(); wb.save(buf); buf.seek(0)
    d1 = await do_test("中文列名", ws)

    # 2. 英文列名
    wb2 = Workbook(); ws2 = wb2.active
    ws2.append(["stock_code", "stock_name", "shares", "cost_price"])
    ws2.append(["600519", "贵州茅台", 100, 1750.00])
    ws2.append(["000858", "五粮液", 200, 135.00])
    buf = io.BytesIO(); wb2.save(buf); buf.seek(0)
    d2 = await do_test("英文列名", ws2)

    # 3. 无列名（内容猜测）
    wb3 = Workbook(); ws3 = wb3.active
    ws3.append(["600519", "贵州茅台", 100, 1750.0])
    ws3.append(["000858", "五粮液", 200, 135.0])
    buf = io.BytesIO(); wb3.save(buf); buf.seek(0)
    d3 = await do_test("无列名(内容猜测)", ws3)

    # 4. 券商变体列名
    wb4 = Workbook(); ws4 = wb4.active
    ws4.append(["证券编码", "股票名称", "持有数量", "成本单价"])
    ws4.append(["600519", "贵州茅台", 100, 1750.00])
    ws4.append(["000858", "五粮液", 200, 135.00])
    buf = io.BytesIO(); wb4.save(buf); buf.seek(0)
    d4 = await do_test("券商变体列名", ws4)

    # Summary
    print("\n" + "="*40)
    passed = [d1,d2,d3,d4]
    for i, d in enumerate(passed, 1):
        tag = "PASS" if d.get("count",0) > 0 else "FAIL"
        print(f"  test{i} count={d.get('count',0)} {tag}")

asyncio.run(test_excel())
